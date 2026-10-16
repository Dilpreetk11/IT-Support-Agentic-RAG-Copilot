"""
Document loading and chunking for the IT Support Copilot ingestion pipeline.

Supports: PDF, TXT, Markdown (.md), DOCX
"""
import logging
from pathlib import Path
from typing import Iterable

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument

from app.core.config import get_settings

logger = logging.getLogger(__name__)
settings = get_settings()

SUPPORTED: set[str] = {".pdf", ".txt", ".md", ".docx"}
MAX_FILE_SIZE_MB = 20
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024


def load_file(path: Path) -> list[Document]:
    """
    Load a document from disk and return a list of LangChain Document objects.

    Raises ValueError for unsupported file types or empty/unreadable documents.
    """
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(
            f"Unsupported file type '{suffix}'. Supported: {', '.join(sorted(SUPPORTED))}"
        )

    # File size check.
    size = path.stat().st_size
    if size > MAX_FILE_SIZE_BYTES:
        raise ValueError(
            f"File too large ({size / 1_048_576:.1f} MB). Maximum allowed: {MAX_FILE_SIZE_MB} MB."
        )
    if size == 0:
        raise ValueError("File is empty.")

    logger.info("Loading document: %s (%d bytes)", path.name, size)

    docs: list[Document] = []

    if suffix == ".pdf":
        docs = PyPDFLoader(str(path)).load()
    elif suffix in {".txt", ".md"}:
        docs = TextLoader(str(path), encoding="utf-8").load()
    elif suffix == ".docx":
        docx = DocxDocument(str(path))
        text = "\n".join(p.text for p in docx.paragraphs if p.text.strip())
        if not text.strip():
            raise ValueError("DOCX file contains no readable text.")
        docs = [Document(page_content=text, metadata={"source": str(path)})]

    # Filter completely empty pages / sections.
    docs = [d for d in docs if d.page_content.strip()]
    if not docs:
        raise ValueError("Document contains no extractable text.")

    logger.info("Loaded %d page(s) from %s.", len(docs), path.name)
    return docs


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    """
    Split documents into overlapping chunks for embedding and retrieval.
    Chunk size and overlap come from application settings.
    """
    doc_list = list(docs)
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        add_start_index=True,
    )
    chunks = splitter.split_documents(doc_list)
    # Remove any chunks that ended up empty after splitting.
    chunks = [c for c in chunks if c.page_content.strip()]
    logger.info("Chunked %d doc(s) into %d non-empty chunks.", len(doc_list), len(chunks))
    return chunks
